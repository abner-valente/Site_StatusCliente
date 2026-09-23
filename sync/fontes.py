"""Camada de entrada: de onde a planilha é lida.

Todo o resto do sistema não sabe qual fonte está em uso. Trocar de fonte é
mudar a variável FONTE_PLANILHA no .env — nenhum outro arquivo muda.

Implementações:
    FonteXlsxLocal    openpyxl, arquivo no disco      — ativa (fase 0)
    FonteExcelOnline  Microsoft Graph                 — esqueleto (fase 3)
    FonteGoogleSheets gspread                         — alternativa não usada
"""

from __future__ import annotations

import os
from typing import Protocol


class FonteDePlanilha(Protocol):
    """Contrato mínimo que o sync exige de qualquer fonte."""

    def ler_aba(self, nome: str) -> list[list[str]]:
        """Devolve a aba inteira como linhas de texto, incluindo o cabeçalho.

        Células vazias viram string vazia, nunca None. Datas viram ISO
        (AAAA-MM-DD) para que o resto do código não precise saber se a fonte
        entrega `datetime` ou texto.
        """
        ...

    def escrever_celula(self, aba: str, linha: int, coluna: str, valor: str) -> None:
        """Grava uma célula. `linha` é 1-based, `coluna` é letra ("Q").

        Usado para uma coisa só: carimbar o uid_royal em linhas novas. Não é
        sincronização de duas vias — é identidade, e nenhum humano digita ali.
        """
        ...


# ---------------------------------------------------------------------------
# Fase 0 — arquivo local
# ---------------------------------------------------------------------------

class FonteXlsxLocal:
    """Lê o .xlsx do disco.

    Serve para desenvolvimento e para a carga inicial. NÃO serve para produção:
    exige alguém rodando o script na própria máquina, e o arquivo fica travado
    enquanto estiver aberto no Excel.
    """

    def __init__(self, caminho: str):
        self.caminho = caminho

    def ler_aba(self, nome: str) -> list[list[str]]:
        import openpyxl

        wb = openpyxl.load_workbook(self.caminho, data_only=True, read_only=True)
        if nome not in wb.sheetnames:
            disponiveis = ", ".join(wb.sheetnames)
            wb.close()
            raise FonteError(
                f"Aba {nome!r} não existe na planilha. Abas disponíveis: {disponiveis}"
            )

        ws = wb[nome]
        linhas = [[_texto(c) for c in linha] for linha in ws.iter_rows(values_only=True)]
        wb.close()
        return linhas

    def escrever_celula(self, aba: str, linha: int, coluna: str, valor: str) -> None:
        import openpyxl

        # read_only=False porque vamos gravar; sem data_only para não destruir
        # as fórmulas das abas Controle Geral e Resumo ao salvar.
        wb = openpyxl.load_workbook(self.caminho)
        wb[aba][f"{coluna}{linha}"] = valor
        wb.save(self.caminho)
        wb.close()


# ---------------------------------------------------------------------------
# Fase 3 — Excel Online (Microsoft Graph)
# ---------------------------------------------------------------------------

class FonteExcelOnline:
    """Lê e escreve a planilha hospedada no OneDrive/SharePoint via Graph.

    NÃO IMPLEMENTADA. O esqueleto abaixo existe para que a troca de fonte seja
    só preencher os métodos, sem mexer em mais nada do sync.

    Pré-requisitos, a levantar com a TI do cliente:

      1. Registro de app no Azure AD com fluxo client credentials
      2. Permissão de aplicativo `Files.ReadWrite.All`
         (ou `Sites.ReadWrite.All` se o arquivo estiver no SharePoint).
         Leitura apenas NÃO basta: o sync escreve o uid_royal de volta.
      3. Consentimento de admin do tenant — costuma ser o gargalo do cronograma
      4. Descomentar msal e requests no requirements.txt

    Esboço da implementação:

        BASE = "https://graph.microsoft.com/v1.0"

        def _token(self) -> str:
            app = msal.ConfidentialClientApplication(
                client_id=self.client_id,
                authority=f"https://login.microsoftonline.com/{self.tenant_id}",
                client_credential=self.client_secret,
            )
            r = app.acquire_token_for_client(
                scopes=["https://graph.microsoft.com/.default"]
            )
            if "access_token" not in r:
                raise FonteError(f"Azure recusou o token: {r.get('error_description')}")
            return r["access_token"]

        def ler_aba(self, nome):
            url = (f"{self.BASE}/drives/{self.drive_id}/items/{self.item_id}"
                   f"/workbook/worksheets('{nome}')/usedRange")
            ...  # resposta traz "values" como lista de listas
            #    normalizar para texto do mesmo jeito que _texto() faz aqui

        def escrever_celula(self, aba, linha, coluna, valor):
            url = (f"{self.BASE}/drives/{self.drive_id}/items/{self.item_id}"
                   f"/workbook/worksheets('{aba}')/range(address='{coluna}{linha}')")
            ...  # PATCH com {"values": [[valor]]}

    Confirmar os caminhos exatos na documentação do Graph na hora de
    implementar — a API muda de versão e não vale confiar em memória.

    Armadilha operacional: o client secret expira (tipicamente 12 ou 24 meses).
    Quando expira, o sync para em silêncio e a planilha continua sendo editada.
    Anote a validade no .env e configure alerta de falha.
    """

    def __init__(self, tenant_id: str, client_id: str, client_secret: str,
                 drive_id: str, item_id: str):
        self.tenant_id = tenant_id
        self.client_id = client_id
        self.client_secret = client_secret
        self.drive_id = drive_id
        self.item_id = item_id

    def ler_aba(self, nome: str) -> list[list[str]]:
        raise NotImplementedError(
            "FonteExcelOnline é da fase 3. Depende do registro de app no Azure AD "
            "com Files.ReadWrite.All e consentimento de admin do tenant."
        )

    def escrever_celula(self, aba: str, linha: int, coluna: str, valor: str) -> None:
        raise NotImplementedError(
            "FonteExcelOnline é da fase 3. Ver o docstring da classe."
        )


# ---------------------------------------------------------------------------
# Apoio
# ---------------------------------------------------------------------------

class FonteError(RuntimeError):
    """Problema ao ler ou escrever na planilha."""


def _texto(valor) -> str:
    """Normaliza uma célula para texto.

    Datas viram ISO para que o resto do sync não precise saber se a fonte
    devolveu `datetime` (openpyxl) ou string (Graph).
    """
    import datetime as _dt

    if valor is None:
        return ""
    if isinstance(valor, _dt.datetime):
        return valor.date().isoformat()
    if isinstance(valor, _dt.date):
        return valor.isoformat()
    return str(valor).strip()


def criar_fonte() -> FonteDePlanilha:
    """Monta a fonte indicada por FONTE_PLANILHA no .env."""
    tipo = os.environ.get("FONTE_PLANILHA", "xlsx_local").strip()

    if tipo == "xlsx_local":
        caminho = os.environ.get("CAMINHO_PLANILHA", "").strip()
        if not caminho:
            raise FonteError("CAMINHO_PLANILHA não definido no .env")
        if not os.path.exists(caminho):
            raise FonteError(f"Planilha não encontrada em {caminho!r}")
        return FonteXlsxLocal(caminho)

    if tipo == "excel_online":
        faltando = [
            v for v in ("AZURE_TENANT_ID", "AZURE_CLIENT_ID", "AZURE_CLIENT_SECRET",
                        "EXCEL_DRIVE_ID", "EXCEL_ITEM_ID")
            if not os.environ.get(v)
        ]
        if faltando:
            raise FonteError(
                "Faltam variáveis para o Excel Online: " + ", ".join(faltando)
            )
        return FonteExcelOnline(
            tenant_id=os.environ["AZURE_TENANT_ID"],
            client_id=os.environ["AZURE_CLIENT_ID"],
            client_secret=os.environ["AZURE_CLIENT_SECRET"],
            drive_id=os.environ["EXCEL_DRIVE_ID"],
            item_id=os.environ["EXCEL_ITEM_ID"],
        )

    raise FonteError(
        f"FONTE_PLANILHA={tipo!r} não reconhecida. Use xlsx_local ou excel_online."
    )
